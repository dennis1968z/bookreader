"""Book Reader – Edge TTS backend (online, low-latency, no local model)."""
import io, re, os, sys, json, asyncio, warnings, base64, uuid
warnings.filterwarnings('ignore')

from flask import Flask, request, jsonify, send_file, Response
import edge_tts
import requests as _http

app = Flask(__name__)


# ─── Supabase helpers ─────────────────────────────────────────────────────────
_SB_URL = os.environ.get('SUPABASE_URL', '').rstrip('/')
_SB_KEY = os.environ.get('SUPABASE_SERVICE_KEY', '')

def _sb_ok():
    return bool(_SB_URL and _SB_KEY)

def _sb_hdrs():
    return {
        'apikey': _SB_KEY,
        'Authorization': f'Bearer {_SB_KEY}',
        'Content-Type': 'application/json',
    }

def sb_save(book_id: str, title: str, cover, paragraphs: list, lang: str) -> bool:
    if not _sb_ok():
        return False
    try:
        r = _http.post(
            f'{_SB_URL}/rest/v1/books',
            headers={**_sb_hdrs(), 'Prefer': 'return=minimal'},
            json={'id': book_id, 'title': title, 'cover': cover,
                  'paragraphs': paragraphs, 'lang': lang},
            timeout=15,
        )
        return r.status_code in (200, 201)
    except Exception:
        return False

def sb_get(book_id: str):
    if not _sb_ok():
        return None
    try:
        r = _http.get(
            f'{_SB_URL}/rest/v1/books?id=eq.{book_id}&select=*',
            headers=_sb_hdrs(),
            timeout=10,
        )
        data = r.json()
        return data[0] if isinstance(data, list) and data else None
    except Exception:
        return None


# ─── Language detection ───────────────────────────────────────────────────────

def detect_lang(text: str) -> str:
    cjk = sum(1 for c in text if '一' <= c <= '鿿')
    return 'z' if cjk / max(len(text), 1) > 0.08 else 'a'


# ─── Text utilities ───────────────────────────────────────────────────────────

def split_paragraphs(text: str) -> list:
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    blocks = re.split(r'\n{2,}', text.strip())
    result = []
    for block in blocks:
        block = re.sub(r'\s+', ' ', block).strip()
        if len(block) < 8:
            continue
        if len(block) > 500:
            sents = re.split(r'(?<=[。！？.!?])\s*', block)
            chunk = ''
            for s in sents:
                if len(chunk) + len(s) <= 450:
                    chunk += s
                else:
                    if chunk.strip():
                        result.append(chunk.strip())
                    chunk = s
            if chunk.strip():
                result.append(chunk.strip())
        else:
            result.append(block)
    return result


# ─── File extractors ──────────────────────────────────────────────────────────

def extract_pdf(data: bytes) -> str:
    try:
        import fitz
        doc = fitz.open(stream=data, filetype='pdf')
        return '\n\n'.join(page.get_text() for page in doc)
    except ImportError:
        pass
    try:
        from pdfminer.high_level import extract_text
        return extract_text(io.BytesIO(data))
    except ImportError:
        raise RuntimeError('PDF needs: pip install pymupdf  or  pip install pdfminer.six')


def extract_epub(data: bytes) -> str:
    try:
        import tempfile, ebooklib
        from ebooklib import epub
        from html.parser import HTMLParser

        class Stripper(HTMLParser):
            def __init__(self):
                super().__init__(); self.parts = []
            def handle_data(self, d):
                if d.strip(): self.parts.append(d)

        with tempfile.NamedTemporaryFile(suffix='.epub', delete=False) as f:
            f.write(data)
            tmp_path = f.name
        try:
            book = epub.read_epub(tmp_path)
        finally:
            os.unlink(tmp_path)
        chapters = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            s = Stripper()
            s.feed(item.get_content().decode('utf-8', errors='replace'))
            chapters.append('\n'.join(s.parts))
        return '\n\n'.join(chapters)
    except ImportError:
        raise RuntimeError('EPUB needs: pip install ebooklib')


def extract_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        return '\n\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        raise RuntimeError('DOCX needs: pip install python-docx')


def extract_doc(data: bytes) -> str:
    """Extract text from old binary .doc (Word 97-2003) via OLE2 parsing."""
    import struct, re

    # Some .doc files are actually OOXML — try python-docx first
    try:
        from docx import Document
        text = '\n\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs if p.text.strip())
        if len(text) > 50:
            return text
    except Exception:
        pass

    # OLE2 binary .doc: scan WordDocument stream for UTF-16LE text runs
    try:
        import olefile
        if not olefile.isOleFile(io.BytesIO(data)):
            raise ValueError('Not OLE2')
        ole = olefile.OleFileIO(io.BytesIO(data))
        try:
            stream = ole.openstream('WordDocument').read()
        finally:
            ole.close()

        parts, run = [], []
        i = 0
        while i < len(stream) - 1:
            try:
                cp = struct.unpack_from('<H', stream, i)[0]
                ch = chr(cp)
                if 0x20 <= cp < 0xD800 and ch.isprintable():
                    run.append(ch)
                elif cp in (0x000D, 0x0007, 0x000C):   # paragraph / page break
                    if run:
                        parts.append(''.join(run)); run = []
                    parts.append('\n')
                elif run:
                    parts.append(''.join(run)); run = []
            except Exception:
                pass
            i += 2
        if run:
            parts.append(''.join(run))

        text = re.sub(r'\n{3,}', '\n\n', ''.join(parts)).strip()
        if len(text) > 50:
            return text
    except ImportError:
        pass
    except Exception:
        pass

    raise RuntimeError('.doc 解析失败，建议在 Word 中另存为 .docx 格式后重试')


# ─── Cover extraction ─────────────────────────────────────────────────────────

def extract_cover(data: bytes, filename: str) -> str | None:
    """Return cover as a base64 data-URL, or None."""
    name = filename.lower()
    try:
        if name.endswith('.pdf'):
            import fitz
            doc = fitz.open(stream=data, filetype='pdf')
            page = doc[0]
            # Scale to ~120px wide thumbnail
            scale = 120 / page.rect.width
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
            raw = pix.tobytes('jpeg')
            if len(raw) > 60_000:          # safety cap
                return None
            return 'data:image/jpeg;base64,' + base64.b64encode(raw).decode()

        if name.endswith('.epub'):
            import tempfile, ebooklib
            from ebooklib import epub
            with tempfile.NamedTemporaryFile(suffix='.epub', delete=False) as f:
                f.write(data); tmp = f.name
            try:
                book = epub.read_epub(tmp)
            finally:
                os.unlink(tmp)
            for item in book.get_items():
                mt = getattr(item, 'media_type', '') or ''
                if not mt.startswith('image/'):
                    continue
                nm = (item.file_name or '').lower()
                if 'cover' in nm or item.get_type() == ebooklib.ITEM_COVER:
                    raw = item.get_content()
                    if len(raw) > 60_000:
                        return None
                    return f'data:{mt};base64,' + base64.b64encode(raw).decode()
    except Exception:
        pass
    return None


# ─── Edge TTS synthesis ───────────────────────────────────────────────────────

async def _synth_async(text: str, voice: str, rate: str):
    communicate = edge_tts.Communicate(text, voice, rate=rate)
    audio = b''
    timings = []
    async for chunk in communicate.stream():
        if chunk['type'] == 'audio':
            audio += chunk['data']
        elif chunk['type'] in ('WordBoundary', 'SentenceBoundary'):
            timings.append({
                'text': chunk['text'],
                'start': round(chunk['offset'] / 10_000_000, 3),
                'end':   round((chunk['offset'] + chunk['duration']) / 10_000_000, 3),
            })
    return audio, timings


def run_async(coro):
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(coro)
    finally:
        loop.close()


# ─── Routes ───────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    return send_file('index.html')


@app.route('/status')
def status():
    return jsonify({'backend': 'edge-tts', 'state': 'ready'})


@app.route('/upload', methods=['POST'])
def upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'No file'}), 400
    name = f.filename.lower()
    data = f.read()
    try:
        if name.endswith(('.txt', '.md', '.text')):
            text = data.decode('utf-8', errors='replace')
        elif name.endswith('.pdf'):
            text = extract_pdf(data)
        elif name.endswith('.epub'):
            text = extract_epub(data)
        elif name.endswith('.docx'):
            text = extract_docx(data)
        elif name.endswith('.doc'):
            text = extract_doc(data)
        else:
            return jsonify({'error': f'Unsupported format: {name.split(".")[-1]}'}), 400
        paras = split_paragraphs(text)
        if not paras:
            return jsonify({'error': 'No text found'}), 400
        lang = detect_lang(' '.join(paras[:10]))
        cover = extract_cover(data, f.filename)
        title = f.filename.rsplit('.', 1)[0] if '.' in f.filename else f.filename
        book_id = str(uuid.uuid4())
        sb_save(book_id, title, cover, paras, lang)
        return jsonify({'paragraphs': paras, 'lang': lang, 'count': len(paras),
                        'cover': cover, 'book_id': book_id})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/book/<book_id>')
def get_book(book_id):
    if not re.match(r'^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$', book_id):
        return jsonify({'error': 'Invalid ID'}), 400
    book = sb_get(book_id)
    if not book:
        return jsonify({'error': 'Book not found'}), 404
    return jsonify(book)


@app.route('/synthesize', methods=['POST'])
def synthesize():
    body = request.json or {}
    text = (body.get('text') or '').strip()
    if not text:
        return jsonify({'error': 'Empty text'}), 400

    lang = body.get('lang') or detect_lang(text)
    default_voice = 'zh-CN-XiaoxiaoNeural' if lang == 'z' else 'en-US-JennyNeural'
    voice = body.get('voice') or default_voice
    speed = float(body.get('speed', 1.0))

    rate_pct = int((speed - 1.0) * 100)
    rate_str = f'+{rate_pct}%' if rate_pct >= 0 else f'{rate_pct}%'

    try:
        audio, timings = run_async(_synth_async(text, voice, rate_str))
    except Exception as e:
        return jsonify({'error': str(e)}), 500

    if not audio:
        return jsonify({'error': 'No audio returned'}), 500

    return Response(
        audio,
        mimetype='audio/mpeg',
        headers={
            'X-Timings': json.dumps(timings, ensure_ascii=True),
            'Access-Control-Expose-Headers': 'X-Timings',
        }
    )


if __name__ == '__main__':
    port = int(os.environ.get('PORT', sys.argv[1] if len(sys.argv) > 1 else 7860))
    print(f'[Book Reader – Edge TTS]  http://localhost:{port}')
    print('  Online TTS via Microsoft Edge – no local model required.')
    print('  Press Ctrl+C to stop.')
    app.run(host='0.0.0.0', port=port, debug=False, threaded=True)
